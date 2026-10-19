"""Adapter for the CBA foreign player registration workbook and birthday-article workbooks."""
from datetime import date, datetime
from pathlib import Path
import re
from openpyxl import load_workbook
from . import foreign

METHOD = 'xlsx merge-header segmentation + season-aware cancellation parsing'


def rows(path, sheet=None):
    book = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        name = sheet or book.sheetnames[0]
        if book.sheetnames != [name]:
            raise ValueError('Foreign source must stay a single-sheet file')
        cells = list(book[name].iter_rows())
        table = [[cell.value for cell in row] for row in cells]
        index = foreign.header_index(table)
        jersey = foreign.columns(table, index)['球衣号码']
        for number in range(index + 2, len(cells)):
            cell = cells[number][jersey]
            value = cell.value
            # A numeric zero displayed with the format 00 is a text code 00.
            # Keep General-format zero as 0, and do not coerce source strings.
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if re.fullmatch(r'0+', cell.number_format or '') and value == int(value):
                    table[number][jersey] = str(int(value)).zfill(len(cell.number_format))
        return name, table
    finally:
        book.close()


def extract(path, season, source, clubs=None, strict=True, sheet=None):
    page, table = rows(path, sheet)
    title = str(table[0][0]) if table and table[0] and table[0][0] else None
    source = {**source, 'method': source.get('method') or METHOD}
    return foreign.build(table, season, source, clubs=clubs, strict=strict,
                         title=title, source_page=page)


# --- Birthday-article workbooks (RC-05) --------------------------------------
# Multi-sheet article workbooks: one article-data sheet plus pivot/summary
# sheets. Pivot/summary/unknown sheets are classification-only and never emit
# fact rows. AC-22: dates are carried verbatim, never interpreted as birthdays
# or publish times. AC-23: nicknames/aliases/titles are retrieval candidates
# only; a Person ID is never assigned by this adapter.
#
# External-source principle (applies to Agent JSON inputs as well):
#   原件保留 -- raw cell values are preserved verbatim in *_raw fields;
#   已知 schema 映射 -- known header synonyms map to canonical columns;
#   未知结构待核验 -- unrecognized layouts classify as unknown (fail-closed),
#   never guessed. This principle does NOT relax any existing validation:
#   Stats-domain checks and fail-closed classification stay as-is.

ARTICLE_DATA = 'article_data'
PIVOT = 'pivot'
SUMMARY = 'summary'
SHEET_UNKNOWN = 'unknown'

ARTICLE_ID_HEADERS = ('文章编号', '微信文章编号', '编号', 'id', 'ID', '文章ID')
ARTICLE_TITLE_HEADERS = ('标题', '文章标题', '题目', 'title')
ARTICLE_DATE_HEADERS = ('日期', 'date', '时间')
# 序号列：识别但不作为 ID（行号语义，非文章身份）。
ARTICLE_SERIAL_HEADERS = ('序号', 'seq', 'No.', 'no')
# 强标记：透视表特有词汇，子串匹配即可（几乎不可能出现在文章明细中）。
STRONG_PIVOT_MARKERS = ('行标签', '数据透视', '透视表', '计数项', '(多项)')
# 弱标记：常见汇总词汇，仅当独立成格或行首出现时才算（避免文章标题含"统计"被误杀）。
WEAK_PIVOT_MARKERS = ('汇总', '总计', '计数', 'Grand Total')
WEAK_SUMMARY_MARKERS = ('统计', '平均', '合计')
# 保留旧名称兼容（内部使用已迁移到强/弱两级）。
PIVOT_MARKERS = STRONG_PIVOT_MARKERS + WEAK_PIVOT_MARKERS
SUMMARY_MARKERS = WEAK_SUMMARY_MARKERS


def _cell_text(value):
    return str(value).strip() if value is not None else ''


def _find_article_header(rows):
    """Shared article-header detector (B2): scan probed text rows 1-12.

    Returns (row_index_0based, id_index, title_index, date_index) or None.
    Used by both _classify_sheet and _article_columns so the two agree on
    which row is the header, including when a title row sits above it.
    """
    for index, cells in enumerate(rows):
        id_index = next((i for i, cell in enumerate(cells) if cell in ARTICLE_ID_HEADERS), None)
        title_index = next((i for i, cell in enumerate(cells) if cell in ARTICLE_TITLE_HEADERS), None)
        if id_index is not None and title_index is not None:
            date_index = next((i for i, cell in enumerate(cells) if cell in ARTICLE_DATE_HEADERS), None)
            return index, id_index, title_index, date_index
    return None


def _is_weak_marker_cell(cell, markers):
    """弱标记仅当独立成格或位于行首时有效（防误杀）。

    文章明细的标题中可能含有"统计"等词（如"数据统计方法"），不能仅因
    子串出现就否决整表。独立成格（cell == marker）或行首出现
    （cell.startswith(marker)）才是汇总行的结构性信号。
    """
    return any(cell == m or cell.startswith(m) for m in markers)


def _find_marker_kind(sheet):
    """Stream the ENTIRE worksheet for pivot/summary markers (B1-R2-01).

    Markers may sit below the header probe window (row 13+); a bounded scan
    would miss them and misclassify a pivot sheet as article data. Any
    marker anywhere vetoes ARTICLE_DATA (fail-closed). Returns PIVOT,
    SUMMARY, or None. Pivot takes precedence, matching _classify_sheet.
    Streams row-by-row on the read-only sheet; early-exits on a pivot hit.

    RC-05 refinement: markers are two-tiered. Strong markers (行标签,
    数据透视, 透视表, 计数项, (多项)) are pivot-specific vocabulary --
    substring match anywhere vetoes. Weak markers (汇总/总计/计数/统计/
    平均/合计) are common words -- they veto only as standalone cells or
    at cell start (structural position), never as substrings inside
    article text (e.g. a title containing "统计" must not kill the sheet).
    """
    found_summary = False
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            cell = _cell_text(value)
            if not cell:
                continue
            if any(marker in cell for marker in STRONG_PIVOT_MARKERS):
                return PIVOT
            if _is_weak_marker_cell(cell, WEAK_PIVOT_MARKERS):
                return PIVOT
            if _is_weak_marker_cell(cell, WEAK_SUMMARY_MARKERS):
                found_summary = True
    return SUMMARY if found_summary else None


def _classify_sheet(sheet):
    # B1-R2-01: the marker scan covers the whole sheet (streaming); the
    # article-header probe stays bounded to rows 1-12 (headers live at top).
    marker = _find_marker_kind(sheet)
    probe = [[_cell_text(value) for value in row]
             for row in sheet.iter_rows(min_row=1, max_row=12, values_only=True)]
    flat = [cell for row in probe for cell in row if cell]
    if not flat and marker is None:
        return SHEET_UNKNOWN
    header = _find_article_header(probe)
    # B1: pivot/summary markers take precedence, fail-closed. A sheet carrying
    # markers is NEVER article data; markers + article-like header coexisting
    # is ambiguous -> unknown (never emitted as facts).
    has_pivot = marker == PIVOT
    has_summary = marker == SUMMARY
    if (has_pivot or has_summary) and header is not None:
        return SHEET_UNKNOWN
    if has_pivot:
        return PIVOT
    if has_summary:
        return SUMMARY
    if header is not None:
        return ARTICLE_DATA
    return SHEET_UNKNOWN


def classify_sheets(path):
    """Classify every sheet of a birthday-article workbook.

    Returns {sheet_name: role} with role in
    {'article_data', 'pivot', 'summary', 'unknown'}. Only 'article_data'
    sheets may be read with article_rows(); all other roles never emit facts.
    """
    book = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        return {name: _classify_sheet(book[name]) for name in book.sheetnames}
    finally:
        book.close()


def _article_columns(sheet):
    probe = [[_cell_text(value) for value in row]
             for row in sheet.iter_rows(min_row=1, max_row=12, values_only=True)]
    found = _find_article_header(probe)
    if found is None:
        raise ValueError('Article header row not found (need id-like and title-like columns)')
    index, id_index, title_index, date_index = found
    return index + 1, id_index, title_index, date_index


def _date_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    # RC-05: Excel 日期序列（如 45705）在 data_only 读取下为 int。
    # 转换为 ISO 文本以便阅读；raw 值原样保留；绝不推断语义（AC-22）。
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 1 <= value <= 80000:  # 约 1900-01-01 至 2118 年，超出则按原文
            try:
                from openpyxl.utils.datetime import from_excel
                dt = from_excel(value)
                return dt.date().isoformat() if hasattr(dt, 'date') else dt.isoformat()
            except (ValueError, OverflowError):
                pass
    return str(value)


def _article_record(values, columns, sheet_name, row_number):
    _, id_index, title_index, date_index = columns

    def at(index):
        return values[index] if index is not None and index < len(values) else None

    raw_id, raw_title, raw_date = at(id_index), at(title_index), at(date_index)
    return {
        'article_id_raw': raw_id,
        'article_id': _cell_text(raw_id) or None,
        'title_raw': raw_title,
        'title': _cell_text(raw_title),
        'row_date_raw': raw_date,
        'row_date_text': _date_text(raw_date),
        # AC-22: dates are carried verbatim; interpretation is never inferred here.
        'date_interpretation': None,
        'date_verified': False,
        # AC-23: nicknames/aliases/titles are retrieval candidates only;
        # a Person ID is never assigned by this adapter.
        'identity_status': 'candidate',
        'person_id': None,
        'source_sheet': sheet_name,
        'source_row': row_number,
    }


def article_rows(path, sheet=None):
    """Read article rows from the data sheet of a birthday-article workbook.

    Returns (sheet_name, [article, ...]). When sheet is omitted, the workbook
    must contain exactly one 'article_data' sheet; otherwise a ValueError is
    raised instead of guessing. Passing a pivot/summary sheet raises
    ValueError: those sheets never emit facts.
    """
    book = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        roles = {name: _classify_sheet(book[name]) for name in book.sheetnames}
        if sheet is not None:
            if sheet not in roles:
                raise ValueError(f'Sheet {sheet!r} not found')
            if roles[sheet] != ARTICLE_DATA:
                raise ValueError(f'Sheet {sheet!r} is {roles[sheet]}, not article data; '
                                 'pivot/summary sheets never emit facts')
            name = sheet
        else:
            data = [sheet_name for sheet_name, role in roles.items() if role == ARTICLE_DATA]
            if len(data) != 1:
                raise ValueError(f'Expected exactly one article-data sheet, found {len(data)}; '
                                 'pass sheet= explicitly')
            name = data[0]
        columns = _article_columns(book[name])
        header_row = columns[0]
        articles = []
        for number, row in enumerate(
                book[name].iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            values = list(row)
            if not any(value is not None and _cell_text(value) for value in values):
                continue
            articles.append(_article_record(values, columns, name, number))
        return name, articles
    finally:
        book.close()
