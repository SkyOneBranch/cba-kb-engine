"""T2-3: birthday-article workbook (.xlsx) extension for foreign_xlsx.

Synthetic fixtures only (generated with openpyxl). Real RC-05 data never
enters the repo. Covers AC-17~AC-23 structural shapes:
  - 89-row article sheet / 89 unique ids / 58 unique titles (synthetic mirror)
  - nickname/alias titles are retrieval candidates, never Person IDs (AC-23)
  - dates are never inferred as birthdays/publish times; raw values kept (AC-22)
  - pivot (Sheet4) and aux summary (Sheet3) sheets are recognized as non-facts
"""
from datetime import datetime
from openpyxl import Workbook
import pytest
from cba_kb.adapters import foreign_xlsx


def birthday_workbook(tmp_path, name='birthday.xlsx'):
    """Synthetic RC-05-shaped workbook: data + pivot + aux summary sheets."""
    book = Workbook()
    data = book.active
    data.title = '工作表1'
    data.append(['文章编号', '标题', '日期'])
    titles = [f'文章标题{i:02d}' for i in range(1, 59)]
    titles[4] = '大狙'      # nickname-like title
    titles[9] = '贺希宁'    # real-name-like title: still must not get a Person ID
    titles[14] = '小卢'     # alias-like title
    for i in range(1, 90):
        title = titles[(i - 1) % 58]
        if i % 3 == 0:
            date_value = f'2025/03/{(i % 28) + 1:02d}'      # string form
        elif i % 3 == 1:
            date_value = None                                # absent
        else:
            date_value = datetime(2025, 3, (i % 28) + 1)     # Excel date form
        data.append([f'A{i:04d}', title, date_value])
    aux = book.create_sheet('Sheet3')
    aux.append(['统计口径说明'])
    aux.append(['文章总数', 89])
    aux.append(['平均每月', 7])
    piv = book.create_sheet('Sheet4')
    piv.append(['数据透视表'])
    piv.append(['标题', '计数'])
    piv.append(['文章标题01', 2])
    path = tmp_path / name
    book.save(path)
    return path


def test_classify_sheets_recognizes_data_pivot_summary(tmp_path):
    roles = foreign_xlsx.classify_sheets(birthday_workbook(tmp_path))
    assert roles == {'工作表1': 'article_data', 'Sheet3': 'summary', 'Sheet4': 'pivot'}


def test_article_rows_structural_shape(tmp_path):
    name, articles = foreign_xlsx.article_rows(birthday_workbook(tmp_path))
    assert name == '工作表1'
    assert len(articles) == 89
    assert len({a['article_id'] for a in articles}) == 89
    assert len({a['title'] for a in articles}) == 58


def test_nickname_titles_are_candidates_never_ids(tmp_path):
    _, articles = foreign_xlsx.article_rows(birthday_workbook(tmp_path))
    assert all(a['person_id'] is None for a in articles)
    assert all(a['identity_status'] == 'candidate' for a in articles)
    by_title = {a['title']: a for a in articles}
    for nickname in ('大狙', '小卢', '贺希宁'):
        row = by_title[nickname]
        assert row['person_id'] is None
        assert row['identity_status'] == 'candidate'
        assert row['title_raw'] == nickname


def test_dates_never_inferred_and_raw_preserved(tmp_path):
    _, articles = foreign_xlsx.article_rows(birthday_workbook(tmp_path))
    assert all(a['date_interpretation'] is None for a in articles)
    assert all(a['date_verified'] is False for a in articles)
    datetimes = [a for a in articles if isinstance(a['row_date_raw'], datetime)]
    strings = [a for a in articles if isinstance(a['row_date_raw'], str)]
    assert datetimes and strings
    for a in datetimes:
        assert a['row_date_text'] == a['row_date_raw'].isoformat()
    for a in strings:
        assert a['row_date_text'] == a['row_date_raw']


def test_pivot_and_summary_sheets_emit_no_facts(tmp_path):
    path = birthday_workbook(tmp_path)
    with pytest.raises(ValueError, match='not article data'):
        foreign_xlsx.article_rows(path, sheet='Sheet4')
    with pytest.raises(ValueError, match='not article data'):
        foreign_xlsx.article_rows(path, sheet='Sheet3')


def test_no_article_data_sheet_is_fail_closed(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet4'
    sheet.append(['数据透视表'])
    sheet.append(['标题', '计数'])
    path = tmp_path / 'nodata.xlsx'
    book.save(path)
    with pytest.raises(ValueError, match='exactly one article-data sheet'):
        foreign_xlsx.article_rows(path)


def test_ambiguous_data_sheets_require_explicit_choice(tmp_path):
    book = Workbook()
    first = book.active
    first.title = '工作表1'
    second = book.create_sheet('工作表2')
    for sheet in (first, second):
        sheet.append(['文章编号', '标题', '日期'])
        sheet.append(['A1', 't', '2025-01-01'])
    path = tmp_path / 'two.xlsx'
    book.save(path)
    with pytest.raises(ValueError, match='exactly one article-data sheet'):
        foreign_xlsx.article_rows(path)
    name, articles = foreign_xlsx.article_rows(path, sheet='工作表2')
    assert name == '工作表2'
    assert len(articles) == 1


def test_source_provenance_points_at_sheet_and_row(tmp_path):
    _, articles = foreign_xlsx.article_rows(birthday_workbook(tmp_path))
    assert articles[0]['source_sheet'] == '工作表1'
    assert articles[0]['source_row'] == 2
    assert articles[-1]['source_row'] == 90


def test_rows_single_sheet_guard_unchanged(tmp_path):
    # Backward compat pin: the foreign-registration path still rejects
    # multi-sheet files; only the new article_* functions handle them.
    with pytest.raises(ValueError, match='single-sheet'):
        foreign_xlsx.rows(birthday_workbook(tmp_path))


# --- Q2-FAIL adversarial regression tests (B1/B2) ---
# Synthetic conflict layouts the original 9 tests did not challenge.


def _conflict_pivot_workbook(tmp_path):
    """Pivot sheet carrying BOTH an article-like header and a pivot marker."""
    book = Workbook()
    sheet = book.active
    sheet.title = '工作表1'
    sheet.append(['文章编号', '标题', '日期'])
    sheet.append(['A0001', '文章标题01', '2025/03/01'])
    sheet.append(['数据透视表：按标题计数'])
    sheet.append(['文章标题01', 2])
    path = tmp_path / 'conflict_pivot.xlsx'
    book.save(path)
    return path


def test_pivot_with_article_like_header_is_fail_closed(tmp_path):
    """B1: pivot marker + article-like header must NEVER emit article facts."""
    path = _conflict_pivot_workbook(tmp_path)
    roles = foreign_xlsx.classify_sheets(path)
    assert roles['工作表1'] != foreign_xlsx.ARTICLE_DATA
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path)
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path, sheet='工作表1')


def _conflict_summary_workbook(tmp_path):
    """Summary sheet carrying BOTH an article-like header and a summary marker."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet3'
    sheet.append(['文章编号', '标题', '日期'])
    sheet.append(['A0001', '文章标题01', '2025/03/01'])
    sheet.append(['合计', 89])
    path = tmp_path / 'conflict_summary.xlsx'
    book.save(path)
    return path


def test_summary_with_article_like_header_is_fail_closed(tmp_path):
    """B1: summary marker + article-like header must NEVER emit article facts."""
    path = _conflict_summary_workbook(tmp_path)
    roles = foreign_xlsx.classify_sheets(path)
    assert roles['Sheet3'] != foreign_xlsx.ARTICLE_DATA
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path, sheet='Sheet3')


def _titled_article_workbook(tmp_path):
    """Article sheet with a title row ABOVE the header row (B2 layout)."""
    book = Workbook()
    sheet = book.active
    sheet.title = '工作表1'
    sheet.append(['球队生日文章清单'])
    sheet.append(['文章编号', '标题', '日期'])
    sheet.append(['A0001', '文章标题01', '2025/03/01'])
    sheet.append(['A0002', '大狙', '2025/03/02'])
    path = tmp_path / 'titled.xlsx'
    book.save(path)
    return path


def test_article_sheet_with_title_row_above_header(tmp_path):
    """B2: header detector is shared — title-above-header works explicitly and auto."""
    path = _titled_article_workbook(tmp_path)
    roles = foreign_xlsx.classify_sheets(path)
    assert roles['工作表1'] == foreign_xlsx.ARTICLE_DATA
    name, articles = foreign_xlsx.article_rows(path, sheet='工作表1')
    assert name == '工作表1'
    assert [a['article_id'] for a in articles] == ['A0001', 'A0002']
    assert articles[0]['source_row'] == 3
    auto_name, auto_articles = foreign_xlsx.article_rows(path)
    assert auto_name == '工作表1'
    assert len(auto_articles) == 2


def _row13_marker_workbook(tmp_path):
    """B1-R2-01: pivot marker BELOW the header probe window (row 13)."""
    book = Workbook()
    sheet = book.active
    sheet.title = '工作表1'
    sheet.append(['文章编号', '标题', '日期'])          # row 1: article-like header
    sheet.append(['A0001', '文章标题01', '2025/03/01'])   # row 2: article row
    for i in range(3, 13):
        sheet.append([f'A{i:04d}', f' filler {i}', '2025/03/01'])  # rows 3-12
    sheet.append(['数据透视表：按标题计数'])  # row 13: pivot marker (probe-invisible)
    sheet.append(['文章标题01', 2])           # row 14: pivot row
    path = tmp_path / 'row13_marker.xlsx'
    book.save(path)
    return path


def test_pivot_marker_on_row13_is_fail_closed(tmp_path):
    """B1-R2-01: a marker below rows 1-12 must still veto ARTICLE_DATA."""
    path = _row13_marker_workbook(tmp_path)
    roles = foreign_xlsx.classify_sheets(path)
    assert roles['工作表1'] != foreign_xlsx.ARTICLE_DATA
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path)
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path, sheet='工作表1')


def test_summary_marker_on_row20_is_fail_closed(tmp_path):
    """B1-R2-01 companion: summary marker below the probe window."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet3'
    sheet.append(['文章编号', '标题', '日期'])
    for i in range(2, 20):
        sheet.append([f'A{i:04d}', f' filler {i}', '2025/03/01'])
    sheet.append(['合计', 89])  # row 20: summary marker
    path = tmp_path / 'row20_summary.xlsx'
    book.save(path)
    roles = foreign_xlsx.classify_sheets(path)
    assert roles['Sheet3'] != foreign_xlsx.ARTICLE_DATA
    with pytest.raises(ValueError):
        foreign_xlsx.article_rows(path, sheet='Sheet3')


def test_full_sheet_marker_scan_perf(tmp_path):
    """B1-R2-01: full-sheet scan stays streaming-fast on a 10k-row sheet."""
    import time
    book = Workbook()
    sheet = book.active
    sheet.title = 'big'
    sheet.append(['文章编号', '标题', '日期'])
    for i in range(2, 10001):
        sheet.append([f'A{i:05d}', f'文章标题{i:05d}', '2025/03/01'])
    path = tmp_path / 'big.xlsx'
    book.save(path)
    start = time.monotonic()
    roles = foreign_xlsx.classify_sheets(path)
    elapsed = time.monotonic() - start
    assert roles['big'] == foreign_xlsx.ARTICLE_DATA
    assert elapsed < 30, f'full-sheet marker scan took {elapsed:.1f}s'
